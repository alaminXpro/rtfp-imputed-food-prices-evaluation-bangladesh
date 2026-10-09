"""BBS monthly CPI loader: national FOOD index and its year-on-year inflation.

Source: Bangladesh Bureau of Statistics (BBS), monthly release "Consumer Price Index (CPI), Inflation Rate and Wage
Rate Index (WRI) in Bangladesh", Table 1.1 "Consumer Price Index (CPI) and Inflation Rate (Point to Point)",
row NATIONAL / Food index and the Inflation row under it. Releases are linked from the BBS "Prices and Wages
(CPI, IIP)" page (PAGE_URL), as PDF (2016 ->) and Excel (November 2024 ->).

Each release prints six monthly columns: the latest three months and the same three months a year earlier. They
are preceded by three fiscal-year annual columns, which this loader drops.

Base. BBS switched the CPI from 2005-06=100 to 2021-22=100 starting with the April 2023 index. The April 2023
release says: "BBS has modernized the CPI compilation from April 2023 ... changing the base index from 2005-06 to
2021-22". In every release this loader has seen, months up to 2023-03 are printed on 2005-06=100 and months from
2023-04 on 2021-22=100. No month is printed on both bases. ``base_for_month`` encodes that rule. The parser then
cross-checks it against the base the page declares, and against the index magnitudes in mixed tables. It raises
instead of guessing. The "Food index" concept also changes with the base: on 2005-06=100 it is the ILO group
"Food, beverage & tobacco"; on 2021-22=100 it is COICOP 01 "Food and non-alcoholic beverages". Levels on
different bases are not comparable, and no splice is made here.

Usage:
    python -I -m rtfp.bbs download [--raw data/external/bbs/raw]
    python -I -m rtfp.bbs build    [--raw data/external/bbs/raw] [--out data/external/bbs]
"""
import argparse
import csv
import datetime as dt
import hashlib
import math
import re
import shutil
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

import numpy as np
import pandas as pd

PAGE_URL = "https://bbs.gov.bd/pages/static-pages/6922de7a933eb65569e1ae8f"
BASE_OLD, BASE_NEW = "2005-06", "2021-22"
BASE_SWITCH_MONTH = "2023-04"  # first month published on 2021-22=100
OBS_COLUMNS = ["source_file", "release_month", "month", "base", "food_index_national",
               "food_inflation_yoy_national"]
MONTHLY_COLUMNS = ["month", "base", "food_index_national", "food_inflation_yoy_national", "source_file"]
MANIFEST_FIELDS = ["filename", "url", "release_month", "sha256"]
PAGE_SNAPSHOT_GLOB = "bbs_prices_wages_page_*.html"

_MONTH_NAMES = ["january", "february", "march", "april", "may", "june", "july", "august", "september",
                "october", "november", "december"]
_MONTH_TYPOS = {"jnly": 7}  # "Jnly'19" in the September 2020 release
_LABEL_RE = re.compile(r"([A-Za-z]{3,9})\s*['’‘`´]\s*(\d{2})(?!\d)")
_NUM_RE = re.compile(r"^-?\d+(?:\.\d+)?$")
_PDF_TOL = 0.005 + 1e-9  # a 2-decimal print vs a full-precision value of the same number
_EXACT_TOL = 1e-9


# --- month labels and base -------------------------------------------------------------------------------------


def _month_number(name: str) -> int | None:
    low = name.lower()
    if low in _MONTH_TYPOS:
        return _MONTH_TYPOS[low]
    hits = [i + 1 for i, full in enumerate(_MONTH_NAMES) if len(low) >= 3 and full.startswith(low)]
    return hits[0] if len(hits) == 1 else None


def parse_month_label(label: str) -> str:
    """BBS column label (e.g. "Jan’22", "Sept'19", "April'15") -> "YYYY-MM". Raises ValueError if unrecognised."""
    m = _LABEL_RE.fullmatch(label.strip())
    num = _month_number(m.group(1)) if m else None
    if num is None:
        raise ValueError(f"not a BBS month label: {label!r}")
    return f"{2000 + int(m.group(2)):04d}-{num:02d}"


def _find_month_labels(text: str) -> list[str]:
    return [m.group(0) for m in _LABEL_RE.finditer(text) if _month_number(m.group(1)) is not None]


def base_for_month(month: str) -> str:
    """CPI base on which BBS publishes a given month ("YYYY-MM")."""
    return BASE_OLD if month < BASE_SWITCH_MONTH else BASE_NEW


def _check_bases(months: list[str], values: list[float], header: str, source_file: str) -> list[str]:
    """Assign bases by the BBS rule and verify them against the table header and the index magnitudes."""
    bases = [base_for_month(m) for m in months]
    # PDF text layers can interleave two header cells, e.g. "(Base Index 2005-06 (Base Index 2021-22 =100) =100)",
    # so a base counts as declared when its label is followed, anywhere later in the header, by "=100".
    flat = re.sub(r"\s+", "", header)
    for base in (BASE_OLD, BASE_NEW):
        if base in bases and not re.search(re.escape(base) + r".*=100", flat):
            raise ValueError(f"{source_file}: has {base}-base months but header does not declare {base}=100")
    old = [v for v, b in zip(values, bases) if b == BASE_OLD]
    new = [v for v, b in zip(values, bases) if b == BASE_NEW]
    if old and new and not min(old) > 2 * max(new):
        raise ValueError(f"{source_file}: index magnitudes contradict the base assignment {list(zip(months, values))}")
    return bases


def _frame(source_file, months, bases, food, yoy) -> pd.DataFrame:
    if len(set(months)) != len(months):
        raise ValueError(f"{source_file}: duplicate month columns {months}")
    release = max(months)
    return pd.DataFrame({"source_file": source_file, "release_month": release, "month": months, "base": bases,
                         "food_index_national": food, "food_inflation_yoy_national": yoy}, columns=OBS_COLUMNS)


# --- PDF text layer --------------------------------------------------------------------------------------------


def _numbers(line: str, label_re: str) -> list[str]:
    rest = re.sub(label_re, "", line, count=1, flags=re.I).split()
    if not all(_NUM_RE.match(t) for t in rest):
        raise ValueError(f"non-numeric token in row: {line!r}")
    return rest


def parse_table_text(text: str, source_file: str) -> pd.DataFrame:
    """Parse Table 1.1 (point to point) text: national food index and YoY inflation for each monthly column."""
    lines = [ln.strip() for ln in text.splitlines()]
    hdr_i = next((i for i, ln in enumerate(lines) if len(_find_month_labels(ln)) >= 3), None)
    if hdr_i is None:
        raise ValueError(f"{source_file}: no month header line")
    months = [parse_month_label(lbl) for lbl in _find_month_labels(lines[hdr_i])]
    nat = next((i for i, ln in enumerate(lines) if ln == "NATIONAL"), None)
    if nat is None:
        raise ValueError(f"{source_file}: no NATIONAL block")
    end = next((i for i in range(nat + 1, len(lines)) if lines[i] in ("RURAL", "URBAN")), len(lines))
    food_i = next((i for i in range(nat + 1, end) if re.match(r"Food\s*index\b", lines[i], re.I)), None)
    if food_i is None:
        raise ValueError(f"{source_file}: no national 'Food index' row")
    food = _numbers(lines[food_i], r"^Food\s*index")
    infl_i = food_i + 1
    if not food:  # numbers printed on the line after the label
        if food_i + 1 >= len(lines):
            raise ValueError(f"{source_file}: national Food index row has no values")
        food, infl_i = _numbers(lines[food_i + 1], r"^"), food_i + 2
    if infl_i >= len(lines) or not re.match(r"Inflation\b", lines[infl_i], re.I):
        raise ValueError(f"{source_file}: no Inflation row after national Food index")
    yoy = _numbers(lines[infl_i], r"^Inflation")
    n = len(months)
    for name, vals in (("food index", food), ("inflation", yoy)):
        if len(vals) != n + 3:
            raise ValueError(f"{source_file}: {name} row has {len(vals)} values, expected {n} monthly + 3 annual")
    food_m = [float(v) for v in food[-n:]]
    yoy_m = [float(v) for v in yoy[-n:]]
    bases = _check_bases(months, food_m, "\n".join(lines[: hdr_i + 1]), source_file)
    return _frame(source_file, months, bases, food_m, yoy_m)


def _is_table_1_1(text: str) -> bool:
    head = re.sub(r"\s+", "", text[:600]).lower()
    return (re.search(r"table:?1(\.1)?\.?consumerpriceindex", head) is not None and "pointtopoint" in head
            and re.search(r"^\s*NATIONAL\s*$", text, re.M) is not None)


def parse_pdf(path: Path) -> pd.DataFrame:
    """Find the Table 1.1 page of a BBS monthly release PDF and parse it."""
    import pdfplumber  # imported here so the Excel path does not need it

    path = Path(path)
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if _is_table_1_1(text):
                return parse_table_text(text, source_file=path.name)
    raise ValueError(f"{path.name}: no Table 1.1 (point to point) page found")


# --- Excel -----------------------------------------------------------------------------------------------------


def parse_xlsx(path: Path) -> pd.DataFrame:
    """Parse sheet "Table-1.1_CPI_Point to Point" of a BBS monthly release workbook."""
    import openpyxl

    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = [ws.title for ws in wb.worksheets
                 if re.search(r"1\.1", ws.title) and re.search(r"point\s*to\s*point", ws.title, re.I)]
        if len(names) != 1:
            raise ValueError(f"{path.name}: expected one Table 1.1 point-to-point sheet, found {names}")
        rows = [list(r) for r in wb[names[0]].iter_rows(values_only=True)]
    finally:
        wb.close()
    first = [str(r[0]).strip() if r and r[0] is not None else "" for r in rows]
    hdr_i, cols = None, []
    for i, r in enumerate(rows):
        cols = [(j, v) for j, v in enumerate(r) if isinstance(v, str) and _find_month_labels(v) == [v.strip()]]
        if len(cols) >= 3:
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError(f"{path.name}: no month header row")
    months = [parse_month_label(v) for _, v in cols]
    nat = next((i for i, v in enumerate(first) if v == "NATIONAL"), None)
    if nat is None:
        raise ValueError(f"{path.name}: no NATIONAL block")
    end = next((i for i in range(nat + 1, len(first)) if first[i] in ("RURAL", "URBAN")), len(first))
    food_i = next((i for i in range(nat + 1, end) if re.match(r"Food\s*index\b", first[i], re.I)), None)
    if food_i is None or food_i + 1 >= len(rows) or not re.match(r"Inflation\b", first[food_i + 1], re.I):
        raise ValueError(f"{path.name}: national Food index / Inflation rows not found")
    vals = []
    for i in (food_i, food_i + 1):
        v = [rows[i][j] for j, _ in cols]
        if not all(isinstance(x, (int, float)) and not isinstance(x, bool) and math.isfinite(x) for x in v):
            raise ValueError(f"{path.name}: non-numeric value in row {i + 1}: {v}")
        vals.append([float(x) for x in v])
    header = "\n".join(" ".join(str(c) for c in r if c is not None) for r in rows[: hdr_i + 1])
    bases = _check_bases(months, vals[0], header, path.name)
    return _frame(path.name, months, bases, vals[0], vals[1])


# --- monthly series --------------------------------------------------------------------------------------------


def _fmt(source_file: str) -> str:
    return source_file.rsplit(".", 1)[-1].lower()


def _is_2dp(x: float) -> bool:
    return abs(x - round(x, 2)) < _EXACT_TOL


def build_monthly(obs: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """One row per month from the latest release that prints it (Excel preferred over PDF of the same release).

    Returns (monthly, revisions). ``revisions`` lists every other print of a month that differs from the chosen
    value beyond rounding: tolerance 0.005 when either value is a 2-decimal print (all PDFs, and the November 2024
    workbook, which stores rounded values), otherwise 1e-9. Two 2-decimal prints that differ by 0.01 are flagged.
    """
    o = obs.copy()
    o["_fmt_rank"] = (o["source_file"].map(_fmt) == "xlsx").astype(int)
    o = o.sort_values(["month", "release_month", "_fmt_rank", "source_file"])
    chosen = o.groupby("month", sort=True).tail(1).set_index("month")
    revs = []
    for _, r in o.iterrows():
        c = chosen.loc[r["month"]]
        if r["source_file"] == c["source_file"]:
            continue
        for field in ("food_index_national", "food_inflation_yoy_national"):
            tol = _PDF_TOL if _is_2dp(r[field]) or _is_2dp(c[field]) else _EXACT_TOL
            if abs(r[field] - c[field]) > tol or r["base"] != c["base"]:
                revs.append({"month": r["month"], "field": field, "chosen_source": c["source_file"],
                             "chosen_value": c[field], "other_source": r["source_file"],
                             "other_release_month": r["release_month"], "other_value": r[field]})
    monthly = chosen.reset_index()[MONTHLY_COLUMNS].reset_index(drop=True)
    cols = ["month", "field", "chosen_source", "chosen_value", "other_source", "other_release_month", "other_value"]
    return monthly, pd.DataFrame(revs, columns=cols)


def within_base_log_change(monthly: pd.DataFrame, k: int = 1) -> pd.Series:
    """ln(I_t / I_{t-k}) of the national food index; NaN where month t-k is absent or on a different base."""
    m = monthly.reset_index(drop=True)
    per = pd.PeriodIndex(m["month"], freq="M")
    idx = {str(p): i for i, p in enumerate(per)}
    out = np.full(len(m), np.nan)
    for i, p in enumerate(per):
        j = idx.get(str(p - k))
        if j is not None and m.at[j, "base"] == m.at[i, "base"]:
            out[i] = math.log(m.at[i, "food_index_national"] / m.at[j, "food_index_national"])
    return pd.Series(out, index=monthly.index, name=f"dlog_food_{k}m")


# --- release page and download ---------------------------------------------------------------------------------


class _PageParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables, self._t, self._row, self._cell, self._a = [], None, None, None, None
        self._heading, self._u = "", None

    def handle_starttag(self, tag, attrs):
        if tag == "u":
            self._u = []
        elif tag == "table":
            self._t = {"heading": self._heading, "rows": []}
        elif tag == "tr" and self._t is not None:
            self._row = []
        elif tag == "td" and self._row is not None:
            self._cell = {"text": [], "links": []}
        elif tag == "a":
            self._a = {"href": dict(attrs).get("href") or "", "text": []}

    def handle_endtag(self, tag):
        if tag == "u" and self._u is not None:
            self._heading, self._u = " ".join("".join(self._u).split()), None
        elif tag == "a" and self._a is not None:
            self._a["text"] = " ".join("".join(self._a["text"]).split())
            if self._cell is not None:
                self._cell["links"].append(self._a)
            self._a = None
        elif tag == "td" and self._cell is not None:
            self._cell["text"] = " ".join("".join(self._cell["text"]).split())
            self._row.append(self._cell)
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self._t["rows"].append(self._row)
            self._row = None
        elif tag == "table" and self._t is not None:
            self.tables.append(self._t)
            self._t = None

    def handle_data(self, data):
        if self._u is not None:
            self._u.append(data)
        if self._a is not None:
            self._a["text"].append(data)
        if self._cell is not None:
            self._cell["text"].append(data)


def _link_format(url: str) -> str:
    return "xlsx" if url.lower().endswith((".xlsx", ".xls")) else "pdf"


def parse_release_page(html: str) -> list[dict]:
    """CPI/WRI monthly-release links on the BBS prices page, labelled with the page's year and month.

    Two layouts are read: blocks of [Year-YYYY | PDF | Excel] (2024 ->) and one column per year with month links
    (2016-2024). The page labels are kept as printed; the release month is verified from file contents later.
    """
    p = _PageParser()
    p.feed(html)
    out = []
    for t in p.tables:
        if not t["heading"].startswith("Monthly Release on Price"):
            continue
        rows = t["rows"]
        hdr = [c["text"] for c in rows[0]]
        if any(re.match(r"Year:\s*\d{4}", h) for h in hdr):
            years = [(re.search(r"\d{4}", h) or [None])[0] for h in hdr]
            for r in rows[1:]:
                for j, c in enumerate(r):
                    for a in c["links"]:
                        out.append({"page_year": years[j], "page_month": a["text"], "format": _link_format(a["href"]),
                                    "url": a["href"]})
        else:
            blocks = [(j, m.group(1)) for j, h in enumerate(hdr) if (m := re.match(r"Year\s*-\s*(\d{4})", h))]
            for r in rows[1:]:
                for j, year in blocks:
                    if j >= len(r) or not r[j]["text"]:
                        continue
                    for k in (1, 2):
                        for a in (r[j + k]["links"] if j + k < len(r) else []):
                            out.append({"page_year": year, "page_month": r[j]["text"],
                                        "format": _link_format(a["href"]), "url": a["href"]})
    return out


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def download(raw: Path, page_url: str = PAGE_URL) -> None:
    """Save a dated snapshot of the BBS page and every CPI monthly-release file it links (TLS verified)."""
    raw.mkdir(parents=True, exist_ok=True)
    with urllib.request.urlopen(page_url, timeout=120) as r:
        html = r.read()
    (raw / f"bbs_prices_wages_page_{dt.date.today().isoformat()}.html").write_bytes(html)
    urls = dict.fromkeys(link["url"] for link in parse_release_page(html.decode("utf-8", "replace")))
    for url in urls:
        dest = raw / url.rsplit("/", 1)[-1]
        if not dest.exists():
            tmp = dest.with_suffix(dest.suffix + ".part")
            with urllib.request.urlopen(url, timeout=300) as r, open(tmp, "wb") as fh:
                shutil.copyfileobj(r, fh)
            tmp.rename(dest)
    print(f"{len(urls)} release files linked; raw dir: {raw}")


# --- build -----------------------------------------------------------------------------------------------------


def parse_all(raw: Path) -> pd.DataFrame:
    frames = []
    for p in sorted(raw.iterdir()):
        if p.suffix.lower() == ".pdf":
            frames.append(parse_pdf(p))
        elif p.suffix.lower() == ".xlsx":
            frames.append(parse_xlsx(p))
    obs = pd.concat(frames, ignore_index=True)
    return obs.sort_values(["release_month", "source_file", "month"]).reset_index(drop=True)


def build_manifest(raw: Path, obs: pd.DataFrame) -> list[dict]:
    snaps = sorted(raw.glob(PAGE_SNAPSHOT_GLOB))
    urls = {}
    if snaps:
        for link in parse_release_page(snaps[-1].read_text(encoding="utf-8", errors="replace")):
            urls.setdefault(link["url"].rsplit("/", 1)[-1], link["url"])
    release = obs.groupby("source_file")["release_month"].first().to_dict()
    rows = []
    for p in sorted(raw.iterdir()):
        if p.name.startswith(".") or p.suffix == ".part":
            continue
        url = PAGE_URL if p.match(PAGE_SNAPSHOT_GLOB) else urls.get(p.name, "")
        rows.append({"filename": p.name, "url": url, "release_month": release.get(p.name, ""),
                     "sha256": sha256_file(p)})
    return rows


def build(raw: Path, out: Path) -> None:
    obs = parse_all(raw)
    monthly, revisions = build_monthly(obs)
    obs.to_csv(out / "bbs_food_cpi_all_releases.csv", index=False)
    monthly.to_csv(out / "bbs_food_cpi_monthly.csv", index=False)
    with open(out / "MANIFEST.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=MANIFEST_FIELDS)
        w.writeheader()
        w.writerows(build_manifest(raw, obs))
    print(f"{obs['source_file'].nunique()} files, {len(obs)} release-month values, {len(monthly)} months "
          f"({monthly['month'].iloc[0]} .. {monthly['month'].iloc[-1]}), {len(revisions)} disagreements")
    if not revisions.empty:
        print(revisions.to_string(index=False))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["download", "build"])
    ap.add_argument("--raw", default="data/external/bbs/raw", type=Path)
    ap.add_argument("--out", default="data/external/bbs", type=Path)
    args = ap.parse_args()
    if args.cmd == "download":
        download(args.raw)
    else:
        build(args.raw, args.out)


if __name__ == "__main__":
    main()
