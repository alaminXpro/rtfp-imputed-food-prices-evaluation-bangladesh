"""Tests for the TCB (Trading Corporation of Bangladesh) daily Dhaka retail price loader.

Fixtures in tests/fixtures/tcb/ are unmodified copies of two files downloaded from the TCB archive
(https://tcb.gov.bd/pages/daily-rmps):
  - 2024_12_d6471fc7...xlsx: the 2021-05-17 bulletin (oldest layout, "Daily Report" sheet)
  - 2026_9_5d9482fc...xlsx:  the 2026-10-08 bulletin (current layout, with a second "changed items" table)
"""
import datetime as dt
import math
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from rtfp.tcb import (
    InvalidTCBFile,
    TCBParseError,
    map_item,
    monthly_summary,
    parse_listing_page,
    parse_many,
    parse_price,
    read_tcb_file,
)

FIX = Path(__file__).parent / "fixtures" / "tcb"
F2021 = FIX / "2024_12_d6471fc763db4e09b8bb9ed6425b6f5d.xlsx"
F2026 = FIX / "2026_9_5d9482fc-f2fa-43e3-8310-98f630aac41c.xlsx"
# TCB states no licence, so fixture files are kept out of git (data/raw/LICENSES.md); skip when absent.
needs_fixtures = pytest.mark.skipif(not (F2021.exists() and F2026.exists()),
                                    reason="TCB fixture files are not redistributed; run locally after download")

TARGETS = {
    ("rice_coarse", "none"),
    ("wheat_flour_atta_loose", "none"),
    ("palm_oil_loose", "none"),
    ("palm_oil_super_loose", "none"),
    ("lentil_masur", "large_grain"),
    ("lentil_masur", "medium_grain"),
    ("lentil_masur", "small_grain"),
}


def _by_key(rows: pd.DataFrame) -> dict:
    return {(r.item_en, r.variant): r for r in rows.itertuples(index=False)}


# --------------------------------------------------------------------------- date from file content

@needs_fixtures
def test_file_date_comes_from_inside_the_file_2021_layout():
    f = read_tcb_file(F2021)
    assert f["file_date"] == dt.date(2021, 5, 17)
    # all three in-file date fields are read and agree
    assert f["date_header"] == dt.date(2021, 5, 17)
    assert f["date_tarikh"] == dt.date(2021, 5, 17)
    assert f["date_text"] == dt.date(2021, 5, 17)
    assert f["date_conflict"] is False


@needs_fixtures
def test_file_date_comes_from_inside_the_file_2026_layout_bengali_digit_tarikh():
    f = read_tcb_file(F2026)
    assert f["file_date"] == dt.date(2026, 10, 8)
    assert f["date_header"] == dt.date(2026, 10, 8)
    assert f["date_tarikh"] == dt.date(2026, 10, 8)  # printed as "তারিখঃ ০৮-১০-২০২৬"
    assert f["date_text"] == dt.date(2026, 10, 8)  # "বৃহস্পতিবার ০৮ অক্টোবর ২০২৬ খ্রিঃ"
    assert f["date_conflict"] is False


@needs_fixtures
def test_file_date_ignores_the_file_name(tmp_path):
    p = tmp_path / "2019-01-01.xlsx"
    p.write_bytes(F2021.read_bytes())
    assert read_tcb_file(p)["file_date"] == dt.date(2021, 5, 17)


def _synthetic(path, header_date, tarikh, text):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ট্রেডিং কর্পোরেশন অব বাংলাদেশ"
    ws["A4"] = text
    ws["A5"] = "স্মারক নং-১"
    ws["K5"] = "তারিখঃ"
    ws["L5"] = tarikh
    ws["A6"], ws["B6"], ws["C6"], ws["E6"] = "পণ্যের নাম ", " মাপের একক", "অদ্যকার মূল্য(টাকায়)", "১সপ্তাহ পূর্বের মূল্য(টাকায়)"
    ws["C7"] = header_date
    ws["C8"], ws["D8"] = "সর্বনিম্ন ", "সর্ব্বোচ্চ "
    ws["A9"], ws["B9"], ws["C9"], ws["D9"] = "চাল (মোটা)/স্বর্ণা/চায়না ইরি", "প্রতি কেজি", "৪৮", "৫২"
    ws["A10"] = "যে সকল বাজার হতে তথ্য সংগ্রহ করা হয়েছেঃ-"
    wb.save(path)
    return path


def test_internal_date_disagreement_is_flagged_and_header_date_wins(tmp_path):
    p = _synthetic(tmp_path / "x.xlsx", dt.datetime(2022, 3, 2), "০১-০৩-২০২২", "বুধবার ০২ মার্চ ২০২২ খ্রিঃ")
    f = read_tcb_file(p)
    assert f["date_header"] == dt.date(2022, 3, 2)
    assert f["date_tarikh"] == dt.date(2022, 3, 1)
    assert f["file_date"] == dt.date(2022, 3, 2)
    assert f["date_conflict"] is True


def test_header_date_as_bengali_text_or_excel_serial(tmp_path):
    p1 = _synthetic(tmp_path / "a.xlsx", "০২-০৩-২০২২", None, "")
    assert read_tcb_file(p1)["file_date"] == dt.date(2022, 3, 2)
    p2 = _synthetic(tmp_path / "b.xlsx", 44622, None, "")  # Excel serial for 2022-03-02
    assert read_tcb_file(p2)["file_date"] == dt.date(2022, 3, 2)



def test_bulletin_on_a_later_sheet_behind_a_decoy_summary_sheet(tmp_path):
    # Some 2022 workbooks carry a TCB summary on Sheet1 (it also has a 'পণ্যের নাম' header but no
    # today's-price column) and the daily bulletin on a later sheet named 'Daily retail price'.
    p = _synthetic(tmp_path / "x.xlsx", dt.datetime(2022, 3, 2), None, "")
    wb = openpyxl.load_workbook(p)
    wb.active.title = "Daily retail price"
    decoy = wb.create_sheet("Sheet1", 0)
    decoy["A2"], decoy["B2"], decoy["C2"] = "ক্রমিক নং", "পণ্যের নাম", "০১-০২-২২ হতে ০১-০৩-২২"
    decoy["B3"], decoy["C3"], decoy["D3"] = "চাল (মোটা)", 1, 2
    wb.save(p)
    f = read_tcb_file(p)
    assert f["sheet"] == "Daily retail price"
    assert f["file_date"] == dt.date(2022, 3, 2)
    r = _by_key(f["rows"])[("rice_coarse", "none")]
    assert (r.min, r.max) == (48.0, 52.0)

# --------------------------------------------------------------------------- item mapping

@pytest.mark.parametrize("label,expected", [
    ("চাল (মোটা)/স্বর্ণা/চায়না ইরি", ("rice_coarse", "none")),
    ("আটা সাদা (খোলা)", ("wheat_flour_atta_loose", "none")),
    ("পাম অয়েল (লুজ)", ("palm_oil_loose", "none")),
    ("পাম অয়েল (সুপার)", ("palm_oil_super_loose", "none")),
    ("সুপার পাম অয়েল (লুজ)", ("palm_oil_super_loose", "none")),
    ("মশুর ডাল (বড় দানা)", ("lentil_masur", "large_grain")),
    ("মশূর ডাল (মাঝারী দানা)", ("lentil_masur", "medium_grain")),
    ("মসুর ডাল (ছোট দানা)", ("lentil_masur", "small_grain")),
    ("  মশুর ডাল   (বড়  দানা) ", ("lentil_masur", "large_grain")),
    ("মশুর\u00a0ডাল (বড়\u200c দানা)", ("lentil_masur", "large_grain")),  # NBSP, zero-width non-joiner
    ("আটা সাদা (খো\u200dলা)", ("wheat_flour_atta_loose", "none")),  # zero-width joiner inside a word
    # look-alikes that must NOT be mapped
    ("চাল (মাঝারী)পাইজাম/আটাশ ", None),
    ("চাল সরু (নাজির/মিনিকেট)", None),
    ("চাল সুগন্ধী (পোলাও)", None),
    ("আটা (প্যাকেট)", None),
    ("ময়দা (খোলা)", None),
    ("আটা/ময়দা", None),
    ("সয়াবিন তেল (লুজ)", None),
    ("রাইস ব্রান তেল (বোতল)", None),
    ("মুগ ডাল (মানভেদে)", None),
    ("এ্যাংকর ডাল", None),
])
def test_map_item(label, expected):
    assert map_item(label) == expected


@needs_fixtures
@pytest.mark.parametrize("path", [F2021, F2026])
def test_fixture_yields_each_target_once(path):
    rows = read_tcb_file(path)["rows"]
    keys = list(zip(rows.item_en, rows.variant))
    assert len(keys) == len(set(keys)), "an item was read twice (e.g. from the 'changed items' table)"
    assert set(keys) == TARGETS


@needs_fixtures
def test_item_bn_is_kept_as_printed():
    rows = _by_key(read_tcb_file(F2021)["rows"])
    assert rows[("lentil_masur", "medium_grain")].item_bn == "মশূর ডাল (মাঝারী দানা)"
    assert rows[("palm_oil_super_loose", "none")].item_bn == "পাম অয়েল (সুপার)"


# --------------------------------------------------------------------------- min / max as numbers

EXPECTED_2021 = {  # values read off the 2021-05-17 bulletin, "today" columns
    ("rice_coarse", "none"): ("প্রতি কেজি", 44, 48),
    ("wheat_flour_atta_loose", "none"): ("প্রতি কেজি", 30, 32),
    ("palm_oil_loose", "none"): ("প্রতি লিটার", 107, 110),
    ("palm_oil_super_loose", "none"): ("প্রতি লিটার", 110, 113),
    ("lentil_masur", "large_grain"): ("প্রতি কেজি", 65, 75),
    ("lentil_masur", "medium_grain"): ("প্রতি কেজি", 80, 90),
    ("lentil_masur", "small_grain"): ("প্রতি কেজি", 100, 110),
}
EXPECTED_2026 = {  # values read off the 2026-10-08 bulletin, "today" columns
    ("rice_coarse", "none"): ("প্রতি কেজি", 50, 55),
    ("wheat_flour_atta_loose", "none"): ("প্রতি কেজি", 50, 50),
    ("palm_oil_loose", "none"): ("প্রতি লিটার", 170, 175),
    ("palm_oil_super_loose", "none"): ("প্রতি লিটার", 175, 180),
    ("lentil_masur", "large_grain"): ("প্রতি কেজি", 90, 100),
    ("lentil_masur", "medium_grain"): ("প্রতি কেজি", 115, 120),
    ("lentil_masur", "small_grain"): ("প্রতি কেজি", 150, 160),
}


@needs_fixtures
@pytest.mark.parametrize("path,expected", [(F2021, EXPECTED_2021), (F2026, EXPECTED_2026)])
def test_min_max_are_numbers_from_todays_columns(path, expected):
    rows = read_tcb_file(path)["rows"]
    assert pd.api.types.is_float_dtype(rows["min"]) and pd.api.types.is_float_dtype(rows["max"])
    got = _by_key(rows)
    for key, (unit, lo, hi) in expected.items():
        assert got[key].unit == unit
        assert got[key].min == lo and got[key].max == hi, key


@pytest.mark.parametrize("raw,expected", [
    (120, 120.0), (120.0, 120.0), ("১২০", 120.0), (" 95 ", 95.0), ("১,২০০", 1200.0), ("৯৫.৫", 95.5),
])
def test_parse_price_numeric(raw, expected):
    assert parse_price(raw) == expected


@pytest.mark.parametrize("raw", [None, "", "  ", "-", "—", 0, "০", "#DIV/0!", "n/a"])
def test_parse_price_missing_or_not_quoted_is_nan(raw):
    assert math.isnan(parse_price(raw))


def test_bengali_digit_prices_in_cells(tmp_path):
    p = _synthetic(tmp_path / "x.xlsx", dt.datetime(2022, 3, 2), None, "")
    r = _by_key(read_tcb_file(p)["rows"])[("rice_coarse", "none")]
    assert (r.min, r.max) == (48.0, 52.0)


# --------------------------------------------------------------------------- invalid files

def test_non_xlsx_bytes_raise_invalid(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"<html><body>404 Not Found</body></html>")
    with pytest.raises(InvalidTCBFile):
        read_tcb_file(p)


def test_empty_file_raises_invalid(tmp_path):
    p = tmp_path / "empty.xlsx"
    p.write_bytes(b"")
    with pytest.raises(InvalidTCBFile):
        read_tcb_file(p)


def test_valid_xlsx_without_price_table_raises_parse_error(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "something else"
    p = tmp_path / "other.xlsx"
    wb.save(p)
    with pytest.raises(TCBParseError):
        read_tcb_file(p)


@needs_fixtures
def test_parse_many_skips_and_lists_invalid_files(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"PK\x03\x04 truncated")
    daily, files, problems = parse_many([F2021, bad, F2026])
    assert [p["filename"] for p in problems] == ["bad.xlsx"]
    assert problems[0]["status"] == "invalid"
    assert set(daily["date"]) == {dt.date(2021, 5, 17), dt.date(2026, 10, 8)}
    assert len(daily) == 2 * len(TARGETS)
    assert list(daily.columns[:7]) == ["date", "item_bn", "item_en", "variant", "unit", "min", "max"]
    assert {f["filename"]: f["file_date"] for f in files} == {
        F2021.name: dt.date(2021, 5, 17), F2026.name: dt.date(2026, 10, 8)}


@needs_fixtures
def test_parse_many_same_content_date_twice_keeps_one_row_when_identical(tmp_path):
    dup = tmp_path / "copy.xlsx"
    dup.write_bytes(F2021.read_bytes())
    daily, _, problems = parse_many([F2021, dup])
    assert problems == []
    assert len(daily) == len(TARGETS)


# --------------------------------------------------------------------------- listing page

LISTING_SNIPPET = """
<table id="noticeTable"><tbody class="table-tbody">
<tr class="table-tr"> <td class="table-td" data-column="id">১</td>
 <!-- <td class="table-td">ঢাকা মহানগরীর নিত্য প্রয়োজনীয় পণ্যের দৈনিক খুচরা বাজার দর</td>--> <!-- <td class="table-td">10/8/2026</td>-->
 <td class="table-td" data-column="title"> ঢাকা মহানগরীর নিত্য প্রয়োজনীয় পণ্যের দৈনিক খুচরা বাজার দর </td>
 <td class="table-td" data-column="publish_time"> 12.30 </td>
 <td class="table-td" data-column="files"> <a class="table-td-icon" href="https://objectstorage.example/o/office-tcb/2026/9/5d9482fc-f2fa-43e3-8310-98f630aac41c.xlsx" target="_blank"> <i></i> </a> </td>
 <td class="table-td" data-column="publish_date"> <span>০৮-১০-২০২৬</span> </td>
 <td class="table-td centered print-none"> <a href="/pages/daily-rmps/x">দেখুন</a> </td> </tr>
<tr class="table-tr"> <td class="table-td" data-column="id">১৪০৮</td>
 <!-- <td class="table-td">t</td>--> <!-- <td class="table-td">1/1/1901</td>-->
 <td class="table-td" data-column="title"> t </td>
 <td class="table-td" data-column="publish_time"> 12.30 </td>
 <td class="table-td" data-column="files"> <a class="table-td-icon" href="https://objectstorage.example/o/office-tcb/2024/12/842169d254ce47aeaaf7a9a555b01590.xlsx"> </a> </td>
 <td class="table-td" data-column="publish_date"> <span></span> </td>
 <td class="table-td centered print-none"> <a href="/pages/daily-rmps/y">দেখুন</a> </td> </tr>
</tbody></table>
<span class="pagination-counts">দেখছেন ১ থেকে ১০ পর্যন্ত, মোট ১৪০৮ এন্ট্রি</span>
"""


def test_parse_listing_page_rows_dates_and_total():
    rows, total = parse_listing_page(LISTING_SNIPPET)
    assert total == 1408
    assert [r["listing_no"] for r in rows] == [1, 1408]
    assert rows[0]["url"].endswith("/2026/9/5d9482fc-f2fa-43e3-8310-98f630aac41c.xlsx")
    assert rows[0]["filename"] == "2026_9_5d9482fc-f2fa-43e3-8310-98f630aac41c.xlsx"
    assert rows[0]["listing_date"] == "2026-10-08"
    assert rows[0]["listing_date_comment"] == "2026-10-08"
    assert rows[1]["listing_date"] == ""  # empty on the site
    assert rows[1]["listing_date_comment"] == "1901-01-01"  # site placeholder, not a real date



# --------------------------------------------------------------------------- manifest: listing vs in-file date

def _entry(no, listing_date, filename, url=None):
    return {"listing_no": no, "title": "t", "publish_time": "12.30", "listing_date": listing_date,
            "listing_date_comment": listing_date, "url": url or f"https://x/o/office-tcb/{filename}",
            "filename": filename}


@needs_fixtures
def test_assemble_keeps_listing_date_but_dates_rows_from_file_content(tmp_path):
    import hashlib

    from rtfp.tcb import assemble

    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / F2021.name).write_bytes(F2021.read_bytes())
    (raw / "copy.xlsx").write_bytes(F2021.read_bytes())  # same bulletin uploaded under a later listing date
    (raw / "zeros.xlsx").write_bytes(b"\x00" * 64)  # like the zero-filled file found in the archive
    (raw / "not_found.jpg").write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF")  # placeholder image
    entries = [_entry(1, "2021-05-17", F2021.name), _entry(2, "2021-05-18", "copy.xlsx"),
               _entry(3, "2023-10-15", "zeros.xlsx"), _entry(4, "", "not_found.jpg")]
    manifest, daily, problems = assemble(entries, raw)
    m = {r["filename"]: r for r in manifest}
    assert m["copy.xlsx"]["listing_date"] == "2021-05-18"
    assert m["copy.xlsx"]["file_date"] == "2021-05-17"
    assert m["copy.xlsx"]["status"] == "ok"
    assert "listing date differs from in-file date" in m["copy.xlsx"]["note"]
    assert f"same bytes as {F2021.name}" in m["copy.xlsx"]["note"]
    assert m[F2021.name]["note"] == ""
    assert m[F2021.name]["sha256"] == hashlib.sha256(F2021.read_bytes()).hexdigest()
    assert m["zeros.xlsx"]["status"] == "invalid" and m["zeros.xlsx"]["file_date"] == ""
    assert m["not_found.jpg"]["status"] == "invalid"
    assert set(daily["date"]) == {dt.date(2021, 5, 17)} and len(daily) == len(TARGETS)
    assert sorted(p["filename"] for p in problems) == ["not_found.jpg", "zeros.xlsx"]

# --------------------------------------------------------------------------- monthly aggregation

def test_monthly_summary_means_and_day_counts():
    d = pd.DataFrame({
        "date": [dt.date(2024, 1, 2), dt.date(2024, 1, 3), dt.date(2024, 2, 1), dt.date(2024, 1, 2)],
        "item_bn": ["a", "a", "a", "b"],
        "item_en": ["lentil_masur"] * 4,
        "variant": ["large_grain", "large_grain", "large_grain", "small_grain"],
        "unit": ["প্রতি কেজি"] * 4,
        "min": [100.0, 110.0, 120.0, 150.0],
        "max": [110.0, 130.0, float("nan"), 160.0],
    })
    m = monthly_summary(d).set_index(["month", "item_en", "variant"])
    jan = m.loc[("2024-01", "lentil_masur", "large_grain")]
    assert jan.n_days == 2
    assert jan.mean_min == pytest.approx(105.0)
    assert jan.mean_max == pytest.approx(120.0)
    assert jan.mean_mid == pytest.approx(((100 + 110) / 2 + (110 + 130) / 2) / 2)
    # variants are not merged
    assert m.loc[("2024-01", "lentil_masur", "small_grain")].n_days == 1
    # a day with a missing max contributes to mean_min only; mid needs both
    feb = m.loc[("2024-02", "lentil_masur", "large_grain")]
    assert feb.n_days == 1 and feb.mean_min == 120.0 and math.isnan(feb.mean_max) and math.isnan(feb.mean_mid)
    assert list(monthly_summary(d).columns) == ["month", "item_en", "variant", "n_days", "mean_min", "mean_max",
                                                "mean_mid"]


# --------------------------------------------------------------------------- TLS

def test_tls_context_pins_missing_intermediate_and_keeps_full_verification():
    import hashlib
    import ssl

    from rtfp.tcb import SECTIGO_DV_R36_PEM, SECTIGO_DV_R36_SHA256, tls_context

    der = ssl.PEM_cert_to_DER_cert(SECTIGO_DV_R36_PEM)
    assert hashlib.sha256(der).hexdigest() == SECTIGO_DV_R36_SHA256
    ctx = tls_context()
    assert ctx.verify_mode == ssl.CERT_REQUIRED and ctx.check_hostname
    # the added intermediate must not become a trust anchor: chains still have to end at a system root
    assert not (ctx.verify_flags & ssl.VERIFY_X509_PARTIAL_CHAIN)
    assert any(c.get("serialNumber") == "397A66CC2756362E0DAA87CA6EABE3B1" for c in ctx.get_ca_certs())
